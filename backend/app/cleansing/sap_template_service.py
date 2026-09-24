"""SAP Template Migration: generates BUT000-General and ADRC-Address rows
from Mandanten, per app/cleansing/sap_template_mappings.py.

Ported from sap_template_module.py's mapping execution engine
(_generate_sheet_frame / _execute_sheet_mappings / _apply_overflow_policies
/ _apply_date_format), simplified to a plain per-row loop over ORM objects
rather than pandas vectorization - this app's Projects are single-tenant
workspaces, not the original's 130k+-row shared database, and every other
cleansing service here already favors a straightforward loop (see
tax_service.py, sap_carp_service.py) over DataFrame vectorization.

Generation is stateless (recomputed fresh from Mandanten/JunkAddress on
every call), matching this app's established pattern for Tax Cleansing's
migration/validation endpoints - no persisted output table to go stale.

Deliberately NOT ported yet (see docs/ROADMAP.md): the four materialized-
table sheets (BUT100/BUT0ID/BUT0IS/BUT000-Append), DFKKBPTAXNUM (needs Tax
Cleansing's TAX_MIGRATION_RESULT - not wired up here), country-filtered
multi-country batch export, CSV output, and anonymization. This slice
proves the mapping engine end-to-end on the two sheets that need neither a
materialized base table nor those extras, and ships one Excel download
covering the whole project.
"""

import re
import uuid

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.cleansing.sap_template_mappings import SHEETS, FieldSpec
from app.models.tenant import JunkAddress, Mandant

_DATE_FMT_TOKEN_RE = re.compile(r"(YYYY|YY|MM|DD|HH|MI|SS)")
_DATE_ISO_RE = re.compile(r"^(\d{4})[-/.]?(\d{2})[-/.]?(\d{2})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?")
_DATE_EU_RE = re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?")


def apply_date_format(value: str, fmt: str) -> tuple[str, bool]:
    """Rewrites a date-ish string into the target token format (YYYY/YY/MM/
    DD/HH/MI/SS, everything else literal). Deliberately not datetime-parsing
    based (fails past year 2262, mis-guesses mixed formats) - regex
    extraction instead (ISO takes precedence over day-first), plus a month
    1-12 / day 1-31 plausibility check. Returns (value, recognized) - an
    empty input returns ("", True); an unrecognized non-empty input returns
    the original value unchanged and False.
    """
    s = (value or "").strip()
    if not s:
        return "", True

    m = _DATE_ISO_RE.match(s)
    if m:
        y, mo, d, h, mi, sec = m.groups()
        mo = (mo or "").zfill(2)
        d = (d or "").zfill(2)
    else:
        m = _DATE_EU_RE.match(s)
        if not m:
            return s, False
        d, mo, y, h, mi, sec = m.groups()
        mo = mo.zfill(2)
        d = d.zfill(2)

    h, mi, sec = h or "00", mi or "00", sec or "00"
    try:
        mo_n, d_n = int(mo), int(d)
    except ValueError:
        return s, False
    if not (1 <= mo_n <= 12) or not (1 <= d_n <= 31):
        return s, False

    result = []
    for seg in _DATE_FMT_TOKEN_RE.split(fmt):
        if not seg:
            continue
        result.append(
            {"YYYY": y, "YY": y[-2:], "MM": mo, "DD": d, "HH": h, "MI": mi, "SS": sec}.get(seg, seg)
        )
    return "".join(result), True


def _get_source_value(mandant: Mandant, field_name: str | None) -> str:
    if not field_name:
        return ""
    return str(getattr(mandant, field_name, None) or "").strip()


def _condition_met(op: str | None, actual: str, expected: str | None) -> bool:
    if op == "not_empty":
        return bool(actual)
    if op == "eq":
        return actual == (expected or "")
    return False


def _apply_field(mandant: Mandant, spec: FieldSpec, join_ids: dict[str, set[str]]) -> tuple[str, bool]:
    """Returns (value, ok) - ok is False only for a FLAG overflow violation."""
    if spec.rule == "NOT_APPLICABLE":
        return "", True

    if spec.rule == "CONSTANT":
        value = spec.fixed_value or ""
    elif spec.rule == "DIRECT_COPY":
        value = _get_source_value(mandant, spec.source)
    elif spec.rule in ("CONDITIONAL_MAP", "CONDITIONAL_COPY"):
        if spec.join_exists_on:
            actual = "x" if mandant.IDParty in join_ids.get(spec.join_exists_on, set()) else ""
        else:
            actual = _get_source_value(mandant, spec.source)
        met = _condition_met(spec.condition_op, actual, spec.condition_value)
        if met and spec.rule == "CONDITIONAL_COPY":
            value = _get_source_value(mandant, spec.then_source)
        elif met:
            value = spec.then_value or ""
        else:
            value = spec.else_value or ""
    else:
        raise ValueError(f"Unsupported rule: {spec.rule}")

    if spec.date_format:
        value, _recognized = apply_date_format(value, spec.date_format)

    ok = True
    if spec.on_overflow == "FLAG" and spec.target_length and len(value) > spec.target_length:
        ok = False
    return value, ok


async def _load_join_ids(db: AsyncSession, project_id: uuid.UUID, models_needed: set[str]) -> dict[str, set[str]]:
    join_ids: dict[str, set[str]] = {}
    if "JunkAddress" in models_needed:
        result = await db.execute(select(JunkAddress.IDParty).where(JunkAddress.project_id == project_id))
        join_ids["JunkAddress"] = {row[0] for row in result.all()}
    return join_ids


async def generate_sheet(
    db: AsyncSession, project_id: uuid.UUID, sheet_name: str, limit: int | None = None
) -> dict:
    """Builds every row of one sheet from this project's Mandanten. Returns
    {"total": n, "rows": [...], "violations": [...]}. `violations` lists
    every FLAG'd overflow, one entry per (row, field).
    """
    if sheet_name not in SHEETS:
        raise ValueError(f"Unknown sheet: {sheet_name}")
    specs = SHEETS[sheet_name]

    result = await db.execute(select(Mandant).where(Mandant.project_id == project_id).order_by(Mandant.IDParty))
    mandanten = list(result.scalars().all())
    total = len(mandanten)
    if limit is not None:
        mandanten = mandanten[:limit]

    join_ids = await _load_join_ids(db, project_id, {s.join_exists_on for s in specs if s.join_exists_on})

    rows: list[dict] = []
    violations: list[dict] = []
    for row_index, m in enumerate(mandanten):
        row: dict[str, str] = {}
        for spec in specs:
            value, ok = _apply_field(m, spec, join_ids)
            row[spec.target_field] = value
            if not ok:
                violations.append(
                    {
                        "row_index": row_index,
                        "id_party": m.IDParty,
                        "field": spec.target_field,
                        "value": value,
                        "allowed_length": spec.target_length,
                    }
                )
        rows.append(row)

    return {"total": total, "rows": rows, "violations": violations}


async def generate_all_sheets(db: AsyncSession, project_id: uuid.UUID) -> dict[str, dict]:
    return {sheet_name: await generate_sheet(db, project_id, sheet_name) for sheet_name in SHEETS}


def build_excel_workbook(sheets_data: dict[str, dict]) -> bytes:
    """One sheet per generated table plus a consolidated "Violations" sheet,
    with every FLAG'd overflow cell highlighted red - mirrors the original's
    write_excel_workbook.
    """
    import io

    import pandas as pd
    from openpyxl.styles import PatternFill

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for sheet_name, data in sheets_data.items():
            columns = list(SHEETS[sheet_name][i].target_field for i in range(len(SHEETS[sheet_name])))
            df = pd.DataFrame(data["rows"], columns=columns)
            # Excel sheet names cap at 31 chars and disallow some punctuation.
            safe_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)
            ws = writer.sheets[safe_name]
            for v in data["violations"]:
                col_idx = columns.index(v["field"]) + 1
                excel_row = v["row_index"] + 2  # header row 1, data starts row 2
                ws.cell(row=excel_row, column=col_idx).fill = red_fill

        all_violations = [
            {"Sheet": sheet_name, **v}
            for sheet_name, data in sheets_data.items()
            for v in data["violations"]
        ]
        pd.DataFrame(
            all_violations, columns=["Sheet", "row_index", "id_party", "field", "value", "allowed_length"]
        ).to_excel(writer, sheet_name="Violations", index=False)

    return buf.getvalue()
